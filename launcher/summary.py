#!/usr/bin/env python3
"""生成总结 Word 文档。

自 2026-09-07 起为「所见即所得」：把小结页面各节内容原样排版成 .docx，
不再经过 AI 改写（AI 改写曾多次丢失子步骤、错置状态；润色保留在页面
「请求润色」按钮，属导出的可选前置步骤，不在此处发生）。
"""

from pathlib import Path


def _markdown_to_docx(markdown_text: str, output_path: str) -> None:
    """Convert simple Markdown to a .docx file (Chinese-friendly).

    Supports #/##/### headings, ``- ``/``* `` bullets, ``1. `` numbered
    lists, and ``| ... |`` pipe tables (separator row dropped, first data
    row bolded as header). 小结页面正文按此规则原样排版，不做内容改写。
    """
    import re

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    # Set default font for the document
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(10.5)
    # Set narrow margins to make better use of the narrow window if printed
    sections = doc.sections[0]
    sections.top_margin = Inches(0.6)
    sections.bottom_margin = Inches(0.6)
    sections.left_margin = Inches(0.7)
    sections.right_margin = Inches(0.7)

    lines = markdown_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line.strip():
            continue
        stripped = line.lstrip()
        if stripped.startswith('# '):
            p = doc.add_heading(stripped[2:], level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith('|'):
            # Collect consecutive pipe-table rows and convert to a Word table.
            rows = [stripped]
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                rows.append(lines[i].lstrip())
                i += 1
            data_rows: list[list[str]] = []
            had_separator = False
            for row in rows:
                cells = [
                    c.strip().replace('**', '')
                    for c in row.strip().strip('|').split('|')
                ]
                if all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                    had_separator = True
                    continue
                data_rows.append(cells)
            if data_rows:
                ncols = max(len(r) for r in data_rows)
                table = doc.add_table(rows=len(data_rows), cols=ncols)
                table.style = 'Table Grid'
                for ri, cells in enumerate(data_rows):
                    for ci in range(ncols):
                        table.cell(ri, ci).text = (
                            cells[ci] if ci < len(cells) else ''
                        )
                # First data row is the header when a separator row existed.
                if had_separator:
                    for ci in range(ncols):
                        runs = table.cell(0, ci).paragraphs[0].runs
                        if runs:
                            runs[0].bold = True
        elif stripped.startswith(('- ', '* ')):
            doc.add_paragraph(stripped[2:], style='List Bullet')
        elif re.match(r'^\d+\.\s', stripped):
            text = re.sub(r'^\d+\.\s', '', stripped)
            doc.add_paragraph(text, style='List Number')
        else:
            doc.add_paragraph(line)

    doc.save(output_path)


def _type_label_map() -> dict[str, tuple[str, str, str]]:
    """Map summary type to (folder_name, doc_name_prefix, period_label)."""
    return {
        'week': ('周报', '周', '周'),
        'month': ('月报', '月', '月'),
        'year': ('年报', '年度', '年'),
    }


def _summary_doc_filename(period_type: str, key: str) -> tuple[str, str]:
    """Return (folder_name, filename.docx) for the given type/key."""
    folder, _, _ = _type_label_map()[period_type]
    if period_type == 'week':
        # key format: "2026-W29"
        year, week = key.split('-W')
        filename = f"{year}年第{int(week)}周工作总结.docx"
    elif period_type == 'month':
        # key format: "2026-07"
        year, month = key.split('-')
        filename = f"{year}年{int(month)}月工作总结.docx"
    else:
        # key format: "2026"
        filename = f"{key}年度工作总结.docx"
    return folder, filename


def generate_summary_doc(
    data_folder_path: str, period_type: str, key: str, sections: dict, ai_config=None
) -> str:
    """把小结合成 Word 文档并保存到共享文件夹。

    所见即所得：docx 与传入 sections 逐节、逐行一致——节标题转二级标题，
    正文行原样转段落，量化汇总表的 ``| ... |`` 行还原为 Word 表格。
    不调用 AI，故不存在句式走样、子步骤丢失或状态错位问题。

    *ai_config* 仅为兼容前端请求保留（历史遗留，不再使用）。

    Returns the saved file path. Raises an exception on failure.
    """
    if not data_folder_path:
        raise Exception("未配置数据文件夹，无法保存总结文档。")

    _ = _type_label_map()[period_type]
    if period_type == 'week':
        year, week = key.split('-W')
        full_label = f"{year}年第{int(week)}周"
    elif period_type == 'month':
        year, month = key.split('-')
        full_label = f"{year}年{int(month)}月"
    else:
        full_label = f"{key}年度"

    # 一级标题 + 每节一个二级标题（空节省略），正文原样保留
    lines = [f"# {full_label}工作总结", ""]
    for title, text in sections.items():
        text = (text or "").strip()
        if not text:
            continue
        lines.append(f"## {title}")
        lines.append(text)
        lines.append("")
    markdown = "\n".join(lines)

    folder_name, filename = _summary_doc_filename(period_type, key)
    output_dir = Path(data_folder_path) / folder_name
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / filename

    _markdown_to_docx(markdown, str(output_path))
    return str(output_path)
