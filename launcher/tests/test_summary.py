from docx import Document

from launcher.summary import (
    _markdown_to_docx,
    _summary_doc_filename,
    generate_summary_doc,
)


def test_summary_doc_filename_week():
    folder, filename = _summary_doc_filename("week", "2026-W29")
    assert folder == "周报"
    assert filename == "2026年第29周工作总结.docx"


def test_summary_doc_filename_month():
    folder, filename = _summary_doc_filename("month", "2026-07")
    assert folder == "月报"
    assert filename == "2026年7月工作总结.docx"


def test_summary_doc_filename_year():
    folder, filename = _summary_doc_filename("year", "2026")
    assert folder == "年报"
    assert filename == "2026年度工作总结.docx"


def test_generate_summary_doc_wysiwyg_matches_sections(tmp_path):
    """导出 docx 与页面小结逐行一致：节标题 + 正文原样，空节省略。"""
    from launcher.state import get_state
    get_state().data_folder_path = str(tmp_path)

    sections = {
        "本周完成任务": (
            "调入人员晋升摸查：已完成，子步骤「摸查更新-调入人员」"
            "「摸查结果上工联单」均已完成。\n"
            "机务跨序列：进行中，4/7 子步骤已完成"
            "（素质测评成绩反馈、过程材料整理）。"
        ),
        "下周计划": "☐ 工班长选配",
        "需协调事项": "",
    }
    path = generate_summary_doc(str(tmp_path), "week", "2026-W29", sections)
    assert path.endswith(".docx")
    docx_path = tmp_path / "周报" / "2026年第29周工作总结.docx"
    assert docx_path.exists()

    doc = Document(str(docx_path))
    paras = [(p.style.name, p.text) for p in doc.paragraphs]
    # 一级标题
    assert paras[0] == ("Heading 1", "2026年第29周工作总结")
    # 各节：二级标题 + 正文行原样保留；空节不出现
    assert ("Heading 2", "本周完成任务") in paras
    assert ("Heading 2", "下周计划") in paras
    assert not any(t == "需协调事项" for _, t in paras)
    for line in sections["本周完成任务"].split("\n"):
        assert ("Normal", line) in paras
    assert ("Normal", "☐ 工班长选配") in paras


def test_generate_summary_doc_month_key(tmp_path):
    """月报 key 解析（2026-07 → 2026年7月工作总结.docx）。"""
    from launcher.state import get_state
    get_state().data_folder_path = str(tmp_path)

    path = generate_summary_doc(
        str(tmp_path), "month", "2026-07", {"本月量化": "本月完成子任务 4 项（跨 3 个任务）"}
    )
    docx_path = tmp_path / "月报" / "2026年7月工作总结.docx"
    assert path.endswith(".docx")
    assert docx_path.exists()
    doc = Document(str(docx_path))
    assert doc.paragraphs[0].text == "2026年7月工作总结"


def test_markdown_to_docx_pipe_table(tmp_path):
    """量化汇总表 `| ... |` 行还原为 Word 表格（分隔行丢弃、表头加粗）。"""
    md = (
        "# 标题\n"
        "\n"
        "## 量化汇总\n"
        "\n"
        "| 分类 | 量化 | 数量 |\n"
        "| --- | --- | --- |\n"
        "| 绩效管理 | 考核 | 120 人 |\n"
        "\n"
        "正文行。"
    )
    out = tmp_path / "table.docx"
    _markdown_to_docx(md, str(out))

    doc = Document(str(out))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert table.cell(0, 0).text == "分类"
    assert table.cell(1, 2).text == "120 人"
    header_run = table.cell(0, 2).paragraphs[0].runs[0]
    assert header_run.bold is True
    assert any(p.text == "正文行。" for p in doc.paragraphs)
